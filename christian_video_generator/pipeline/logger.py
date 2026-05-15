"""
Run logger — appends one row per video to output/videos_log.xlsx.

Never overwrites existing rows.  Creates the file and headers on first use.
"""

from datetime import datetime
from pathlib import Path

_LOG_FILE = Path(__file__).parent.parent / "output" / "videos_log.xlsx"
_COLUMNS = ["Date", "Theme", "Title", "Output File", "Duration (s)", "Segments", "Status"]


def _get_workbook():
    """Load existing workbook or create a new one with headers."""
    try:
        import openpyxl
    except ImportError:
        raise ImportError(
            "openpyxl is required for logging. "
            "Install it with: pip install openpyxl"
        )

    _LOG_FILE.parent.mkdir(parents=True, exist_ok=True)

    if _LOG_FILE.exists():
        wb = openpyxl.load_workbook(_LOG_FILE)
        ws = wb.active
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Video Log"
        ws.append(_COLUMNS)
        # Basic header styling
        from openpyxl.styles import Font, PatternFill
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
        # Column widths
        widths = [16, 30, 35, 45, 14, 10, 12]
        for col, width in zip(ws.iter_cols(min_row=1, max_row=1), widths):
            ws.column_dimensions[col[0].column_letter].width = width

    return wb, ws


def log_run(
    theme: str,
    title: str,
    output_path: Path,
    duration: float,
    segments_count: int,
    status: str = "success",
    timestamp: datetime | None = None,
) -> None:
    """
    Append one log row to output/videos_log.xlsx.

    Parameters
    ----------
    theme          : the --theme CLI argument
    title          : AI-generated video title
    output_path    : final MP4 path
    duration       : video duration in seconds
    segments_count : number of narration segments assembled
    status         : 'success' or 'failed'
    timestamp      : defaults to now()
    """
    if timestamp is None:
        timestamp = datetime.now()

    try:
        wb, ws = _get_workbook()
        ws.append([
            timestamp.strftime("%Y-%m-%d %H:%M:%S"),
            theme,
            title,
            str(output_path),
            round(duration, 1),
            segments_count,
            status,
        ])
        wb.save(_LOG_FILE)
        print(f"  [logger] run logged → {_LOG_FILE}")
    except Exception as exc:
        # Logging failure must never abort the pipeline
        print(f"  [logger] WARNING: could not write log: {exc}")
